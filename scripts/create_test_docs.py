#!/usr/bin/env python3
"""
Test Document Generator for IlluminAI

This script creates test documents in various formats supported by IlluminAI.
Run this script from the project root directory to generate sample documents
in the test_documents folder.
"""

import os
import sys
import docx
import markdown
from PyPDF2 import PdfWriter, PdfReader
from io import BytesIO
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter

# Sample text content
SAMPLE_TEXT = """# Sample Test Document

This is a test document for IlluminAI Document Reader.

## Testing Features

You can use this document to test:
* Text selection
* AI explanations
* Summarization
* Question answering

## Sample Content

### What is Artificial Intelligence?

Artificial Intelligence (AI) refers to the simulation of human intelligence in machines. 
These systems can perform tasks that typically require human intelligence, such as:
- Visual perception
- Speech recognition
- Decision-making
- Language translation

### Example Code

```python
def hello_illuminai():
    print("Hello from IlluminAI!")
    return "Testing document interaction"
```

### Mathematical Formula

The formula for calculating the area of a circle is:
A = πr²

Where:
- A is the area
- r is the radius
- π is approximately 3.14159

## End of Document

This is the end of the test document. Thank you for trying IlluminAI!
"""


def ensure_test_dir():
    """Ensure the test_documents directory exists."""
    test_dir = "test_documents"
    if not os.path.exists(test_dir):
        os.makedirs(test_dir)
    return test_dir


def create_text_file(test_dir):
    """Create a sample .txt file."""
    file_path = os.path.join(test_dir, "sample.txt")
    with open(file_path, "w") as f:
        f.write(SAMPLE_TEXT)
    print(f"Created text file: {file_path}")


def create_markdown_file(test_dir):
    """Create a sample .md file."""
    file_path = os.path.join(test_dir, "sample.md")
    with open(file_path, "w") as f:
        f.write(SAMPLE_TEXT)
    print(f"Created markdown file: {file_path}")


def create_docx_file(test_dir):
    """Create a sample .docx file."""
    file_path = os.path.join(test_dir, "sample.docx")

    doc = docx.Document()
    doc.add_heading("Sample Test Document", 0)
    doc.add_paragraph("This is a test document for IlluminAI Document Reader.")

    doc.add_heading("Testing Features", 1)
    features = doc.add_paragraph()
    features.add_run("You can use this document to test:\n")
    doc.add_paragraph("Text selection", style="List Bullet")
    doc.add_paragraph("AI explanations", style="List Bullet")
    doc.add_paragraph("Summarization", style="List Bullet")
    doc.add_paragraph("Question answering", style="List Bullet")

    doc.add_heading("What is Artificial Intelligence?", 2)
    doc.add_paragraph(
        "Artificial Intelligence (AI) refers to the simulation of human "
        "intelligence in machines. These systems can perform tasks that "
        "typically require human intelligence."
    )

    tasks = doc.add_paragraph()
    tasks.add_run("Common AI tasks include:\n")
    doc.add_paragraph("Visual perception", style="List Bullet")
    doc.add_paragraph("Speech recognition", style="List Bullet")
    doc.add_paragraph("Decision-making", style="List Bullet")
    doc.add_paragraph("Language translation", style="List Bullet")

    doc.add_heading("Example Code", 2)
    doc.add_paragraph(
        "def hello_illuminai():\n"
        "    print(\"Hello from IlluminAI!\")\n"
        "    return \"Testing document interaction\""
    )

    doc.add_heading("Mathematical Formula", 2)
    doc.add_paragraph(
        "The formula for calculating the area of a circle is:\n"
        "A = πr²\n\n"
        "Where:\n"
        "- A is the area\n"
        "- r is the radius\n"
        "- π is approximately 3.14159"
    )

    doc.add_heading("End of Document", 1)
    doc.add_paragraph(
        "This is the end of the test document. Thank you for trying IlluminAI!"
    )

    doc.save(file_path)
    print(f"Created Word document: {file_path}")


def create_pdf_file(test_dir):
    """Create a sample .pdf file."""
    file_path = os.path.join(test_dir, "sample.pdf")

    # Create PDF content using ReportLab
    buffer = BytesIO()
    c = canvas.Canvas(buffer, pagesize=letter)
    width, height = letter

    # Title
    c.setFont("Helvetica-Bold", 16)
    c.drawString(72, height - 72, "Sample Test Document")

    # Introduction
    c.setFont("Helvetica", 12)
    c.drawString(72, height - 100,
                 "This is a test document for IlluminAI Document Reader.")

    # Testing Features
    c.setFont("Helvetica-Bold", 14)
    c.drawString(72, height - 130, "Testing Features")

    c.setFont("Helvetica", 12)
    c.drawString(72, height - 150, "You can use this document to test:")
    c.drawString(90, height - 170, "• Text selection")
    c.drawString(90, height - 190, "• AI explanations")
    c.drawString(90, height - 210, "• Summarization")
    c.drawString(90, height - 230, "• Question answering")

    # What is AI
    c.setFont("Helvetica-Bold", 14)
    c.drawString(72, height - 260, "What is Artificial Intelligence?")

    c.setFont("Helvetica", 12)
    text = "Artificial Intelligence (AI) refers to the simulation of human intelligence in machines."
    c.drawString(72, height - 280, text)
    text = "These systems can perform tasks that typically require human intelligence, such as:"
    c.drawString(72, height - 300, text)

    c.drawString(90, height - 320, "• Visual perception")
    c.drawString(90, height - 340, "• Speech recognition")
    c.drawString(90, height - 360, "• Decision-making")
    c.drawString(90, height - 380, "• Language translation")

    # Example Code
    c.setFont("Helvetica-Bold", 14)
    c.drawString(72, height - 410, "Example Code")

    c.setFont("Courier", 12)
    c.drawString(72, height - 430, "def hello_illuminai():")
    c.drawString(90, height - 450, "print(\"Hello from IlluminAI!\")")
    c.drawString(90, height - 470, "return \"Testing document interaction\"")

    # Mathematical Formula
    c.setFont("Helvetica-Bold", 14)
    c.drawString(72, height - 500, "Mathematical Formula")

    c.setFont("Helvetica", 12)
    c.drawString(72, height - 520,
                 "The formula for calculating the area of a circle is:")
    c.drawString(72, height - 540, "A = πr²")

    c.drawString(72, height - 560, "Where:")
    c.drawString(90, height - 580, "• A is the area")
    c.drawString(90, height - 600, "• r is the radius")
    c.drawString(90, height - 620, "• π is approximately 3.14159")

    # End of Document
    c.setFont("Helvetica-Bold", 14)
    c.drawString(72, height - 650, "End of Document")

    c.setFont("Helvetica", 12)
    c.drawString(72, height - 670,
                 "This is the end of the test document. Thank you for trying IlluminAI!")

    c.save()

    # Create PDF file
    buffer.seek(0)
    with open(file_path, "wb") as f:
        f.write(buffer.getvalue())

    print(f"Created PDF document: {file_path}")


def create_rtf_file(test_dir):
    """Create a sample .rtf file using a simple RTF template."""
    file_path = os.path.join(test_dir, "sample.rtf")

    # Simple RTF header and formatting
    rtf_content = """{\\rtf1\\ansi\\ansicpg1252\\cocoartf2639
\\cocoatextscaling0\\cocoaplatform0{\\fonttbl\\f0\\fswiss\\fcharset0 Helvetica-Bold;\\f1\\fswiss\\fcharset0 Helvetica;}
{\\colortbl;\\red255\\green255\\blue255;}
{\\*\\expandedcolortbl;;}
\\margl1440\\margr1440\\vieww11520\\viewh8400\\viewkind0
\\pard\\tx720\\tx1440\\tx2160\\tx2880\\tx3600\\tx4320\\tx5040\\tx5760\\tx6480\\tx7200\\tx7920\\tx8640\\pardirnatural\\partightenfactor0

\\f0\\b\\fs36 \\cf0 Sample Test Document
\\f1\\b0\\fs24 \\
\\
This is a test document for IlluminAI Document Reader.\\
\\

\\f0\\b\\fs28 Testing Features
\\f1\\b0\\fs24 \\
\\
You can use this document to test:\\
\\pard\\tx220\\tx720\\tx1440\\tx2160\\tx2880\\tx3600\\tx4320\\tx5040\\tx5760\\tx6480\\tx7200\\tx7920\\tx8640\\li720\\fi-720\\pardirnatural\\partightenfactor0
\\ls1\\ilvl0\\cf0 {\\listtext	\\uc0\\u8226 	}Text selection\\
{\\listtext	\\uc0\\u8226 	}AI explanations\\
{\\listtext	\\uc0\\u8226 	}Summarization\\
{\\listtext	\\uc0\\u8226 	}Question answering\\
\\pard\\tx720\\tx1440\\tx2160\\tx2880\\tx3600\\tx4320\\tx5040\\tx5760\\tx6480\\tx7200\\tx7920\\tx8640\\pardirnatural\\partightenfactor0
\\cf0 \\

\\f0\\b\\fs28 What is Artificial Intelligence?
\\f1\\b0\\fs24 \\
\\
Artificial Intelligence (AI) refers to the simulation of human intelligence in machines. These systems can perform tasks that typically require human intelligence, such as:\\
\\pard\\tx220\\tx720\\tx1440\\tx2160\\tx2880\\tx3600\\tx4320\\tx5040\\tx5760\\tx6480\\tx7200\\tx7920\\tx8640\\li720\\fi-720\\pardirnatural\\partightenfactor0
\\ls2\\ilvl0\\cf0 {\\listtext	\\uc0\\u8226 	}Visual perception\\
{\\listtext	\\uc0\\u8226 	}Speech recognition\\
{\\listtext	\\uc0\\u8226 	}Decision-making\\
{\\listtext	\\uc0\\u8226 	}Language translation\\
\\pard\\tx720\\tx1440\\tx2160\\tx2880\\tx3600\\tx4320\\tx5040\\tx5760\\tx6480\\tx7200\\tx7920\\tx8640\\pardirnatural\\partightenfactor0
\\cf0 \\

\\f0\\b\\fs28 Example Code
\\f1\\b0\\fs24 \\
\\
def hello_illuminai():\\
    print("Hello from IlluminAI!")\\
    return "Testing document interaction"\\
\\

\\f0\\b\\fs28 Mathematical Formula
\\f1\\b0\\fs24 \\
\\
The formula for calculating the area of a circle is:\\
A = \\uc0\\u960 r\\uc0\\u178 \\
\\
Where:\\
\\pard\\tx220\\tx720\\tx1440\\tx2160\\tx2880\\tx3600\\tx4320\\tx5040\\tx5760\\tx6480\\tx7200\\tx7920\\tx8640\\li720\\fi-720\\pardirnatural\\partightenfactor0
\\ls3\\ilvl0\\cf0 {\\listtext	\\uc0\\u8226 	}A is the area\\
{\\listtext	\\uc0\\u8226 	}r is the radius\\
{\\listtext	\\uc0\\u8226 	}\\uc0\\u960  is approximately 3.14159\\
\\pard\\tx720\\tx1440\\tx2160\\tx2880\\tx3600\\tx4320\\tx5040\\tx5760\\tx6480\\tx7200\\tx7920\\tx8640\\pardirnatural\\partightenfactor0
\\cf0 \\

\\f0\\b\\fs28 End of Document
\\f1\\b0\\fs24 \\
\\
This is the end of the test document. Thank you for trying IlluminAI!}"""

    with open(file_path, "w") as f:
        f.write(rtf_content)

    print(f"Created RTF document: {file_path}")


def main():
    """Main function to create test documents."""
    print("Creating test documents for IlluminAI...")

    try:
        test_dir = ensure_test_dir()

        create_text_file(test_dir)
        create_markdown_file(test_dir)

        try:
            create_docx_file(test_dir)
        except Exception as e:
            print(f"Warning: Could not create DOCX file: {e}")

        try:
            create_pdf_file(test_dir)
        except Exception as e:
            print(f"Warning: Could not create PDF file: {e}")

        create_rtf_file(test_dir)

        print("\nTest documents created successfully in the 'test_documents' directory.")
        print("You can now open these documents with IlluminAI Document Reader.")

    except Exception as e:
        print(f"Error creating test documents: {e}")
        return 1

    return 0


if __name__ == "__main__":
    sys.exit(main())
